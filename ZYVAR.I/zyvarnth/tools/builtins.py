from __future__ import annotations

import base64
import hashlib
from io import BytesIO
import os
import platform
import shutil
import stat
import subprocess
import webbrowser
import json
from urllib.parse import urlencode
from urllib.request import Request, urlopen
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from ..config import Settings
from ..database import Database
from ..security import approved_path, redact_sensitive_text, safe_pattern, validate_http_url
from .browser import BrowserController
from .registry import ApprovalRequired, ToolRegistry, ToolSpec
from .system_adapters import SystemSettingsAdapter

MAX_TEXT_FILE = 1_000_000
MAX_DOCUMENT_BYTES = 20_000_000
_WEATHER_DESCRIPTIONS = {
    0: "Clear sky", 1: "Mainly clear", 2: "Partly cloudy", 3: "Overcast",
    45: "Fog", 48: "Depositing rime fog", 51: "Light drizzle", 53: "Moderate drizzle",
    55: "Dense drizzle", 56: "Light freezing drizzle", 57: "Dense freezing drizzle",
    61: "Slight rain", 63: "Moderate rain", 65: "Heavy rain", 66: "Light freezing rain",
    67: "Heavy freezing rain", 71: "Slight snowfall", 73: "Moderate snowfall",
    75: "Heavy snowfall", 77: "Snow grains", 80: "Slight rain showers",
    81: "Moderate rain showers", 82: "Violent rain showers", 85: "Slight snow showers",
    86: "Heavy snow showers", 95: "Thunderstorm", 96: "Thunderstorm with slight hail",
    97: "Heavy thunderstorm", 99: "Thunderstorm with heavy hail",
}
_APP_COMMANDS = {
    "calculator": {"Darwin": ["open", "-a", "Calculator"], "Windows": ["calc.exe"], "Linux": ["gnome-calculator"]},
    "notes": {"Darwin": ["open", "-a", "Notes"], "Windows": ["notepad.exe"], "Linux": ["gedit"]},
    "browser": {"Darwin": ["open", "-a", "Safari"], "Windows": ["cmd", "/c", "start", "", "about:blank"], "Linux": ["xdg-open", "about:blank"]},
}


class ToolExecutor:
    def __init__(self, settings: Settings, database: Database) -> None:
        self.settings = settings
        self.db = database
        self.registry = ToolRegistry()
        self._install_builtins()

    @property
    def approved_roots(self) -> list[str]:
        return list(self.settings.get("approved_roots", []))

    def _install_builtins(self) -> None:
        add = self.registry.register
        add(ToolSpec("open_url", "Open a web address in the user's default browser.", {"url": (str, True)}, self.open_url))
        add(ToolSpec("open_application", "Open one of the supported applications: calculator, notes, or browser.", {"name": (str, True)}, self.open_application))
        add(ToolSpec("system_status", "Read CPU, memory, and battery telemetry available on this device.", {}, self.system_status))
        add(ToolSpec("set_volume", "Set the device output volume, if the operating system supports it.", {"percent": (int, True)}, self.set_volume, False, True))
        add(ToolSpec("web_search", "Search the live web and return source titles, URLs, and snippets.", {"query": (str, True), "limit": (int, False)}, self.web_search))
        add(ToolSpec("get_weather", "Fetch current weather from Open-Meteo for a named location.", {"location": (str, True)}, self.get_weather))
        add(ToolSpec("browser_visit", "Load a web page with Playwright and return its visible text.", {"url": (str, True)}, self.browser_visit))
        add(ToolSpec("browser_interact", "In an isolated Playwright page, fill one field or click one element. A click may submit data.", {"url": (str, True), "action": (str, True), "selector": (str, True), "value": (str, False)}, self.browser_interact, True))
        add(ToolSpec("file_search", "Search filenames inside folders the user approved.", {"pattern": (str, True)}, self.file_search))
        add(ToolSpec("file_read", "Read a text file inside a folder the user approved.", {"path": (str, True)}, self.file_read))
        add(ToolSpec("file_write", "Create or replace a text file inside an approved folder.", {"path": (str, True), "content": (str, True)}, self.file_write, True, True))
        add(ToolSpec("file_move", "Move a file between approved folders without overwriting an existing file.", {"source": (str, True), "destination": (str, True)}, self.file_move, True, True))
        add(ToolSpec("file_delete", "Move an approved file to Zyvarnth's recoverable trash.", {"path": (str, True)}, self.file_delete, True, True))
        add(ToolSpec("desktop_type_text", "Type text into the currently focused application.", {"text": (str, True)}, self.desktop_type_text, True))
        add(ToolSpec("desktop_shortcut", "Press a validated keyboard shortcut in the focused application.", {"keys": (list, True)}, self.desktop_shortcut, True))
        add(ToolSpec("create_reminder", "Create an in-app reminder in the next 365 days.", {"title": (str, True), "minutes": (int, True)}, self.create_reminder))

    def execute(self, name: str, arguments: dict[str, Any], *, approved: bool = False) -> Any:
        spec = self.registry.get(name)
        if spec is None:
            raise PermissionError(f"Unknown or unauthorized tool: {name}")
        self._validate(spec, arguments)
        if spec.requires_confirmation and not approved:
            raise ApprovalRequired(name, self.describe(spec, arguments))
        try:
            result = spec.handler(**arguments)
        except Exception as exc:
            # Arguments are intentionally excluded from failure journaling.
            error = redact_sensitive_text(str(exc))
            self.db.add_action(name, f"Failed: {error[:220]}", {}, False)
            raise
        reversible = spec.reversible
        if name == "set_volume":
            reversible = (
                isinstance(result, dict)
                and result.get("before") is not None
                and result.get("after") == arguments.get("percent")
            )
        payload = result if reversible and isinstance(result, dict) else {}
        summary = redact_sensitive_text(self.log_summary(spec, arguments))
        if name == "set_volume" and not reversible:
            summary += " (change applied but undo could not be verified)"
        self.db.add_action(name, summary, payload, reversible)
        return result

    @staticmethod
    def _validate(spec: ToolSpec, arguments: dict[str, Any]) -> None:
        if not isinstance(arguments, dict):
            raise ValueError("Tool arguments must be an object.")
        allowed = set(spec.parameters)
        if set(arguments) - allowed:
            raise ValueError("Unexpected tool argument supplied.")
        for key, (kind, required) in spec.parameters.items():
            if required and key not in arguments:
                raise ValueError(f"Missing required argument: {key}")
            if key in arguments and (not isinstance(arguments[key], kind) or (kind is int and isinstance(arguments[key], bool))):
                raise ValueError(f"Invalid value for {key}.")

    @staticmethod
    def describe(spec: ToolSpec, args: dict[str, Any]) -> str:
        if spec.name == "file_write":
            preview = args["content"][:240]
            return f"Write file: {args['path']} ({len(args['content'])} characters)\nPreview: {preview}"
        if spec.name == "desktop_type_text":
            return f"Type this into the focused app ({len(args['text'])} characters):\n{args['text'][:500]}"
        if spec.name == "browser_interact":
            risk = " A click may submit a form or send data." if args.get("action") == "click" else ""
            value = f"\nText to fill: {args.get('value', '')[:500]}" if args.get("action") == "fill" else ""
            return f"Browser {args.get('action')} at {args.get('url')} using selector {args.get('selector')}.{risk}{value}"
        if spec.name == "file_delete":
            return f"Move this file to recoverable trash in the same folder: {args['path']}"
        return f"{spec.name.replace('_', ' ').title()}: " + ", ".join(f"{key}={value}" for key, value in args.items())

    @staticmethod
    def log_summary(spec: ToolSpec, args: dict[str, Any]) -> str:
        if spec.name == "file_write":
            return f"Write file: {args['path']} ({len(args['content'])} characters)"
        if spec.name == "desktop_type_text":
            return f"Type {len(args['text'])} characters into the focused app"
        if spec.name == "browser_interact":
            return f"Browser {args.get('action')} at {args.get('url')} using selector {args.get('selector')}"
        return ToolExecutor.describe(spec, args)

    def open_url(self, url: str) -> dict[str, Any]:
        safe = validate_http_url(url)
        if not webbrowser.open(safe, new=2):
            raise RuntimeError("The system did not accept the browser launch request.")
        return {"opened": safe}

    def open_application(self, name: str) -> dict[str, str]:
        key = name.strip().lower()
        commands = _APP_COMMANDS.get(key)
        if not commands:
            raise ValueError("Supported app names are calculator, notes, and browser.")
        command = commands.get(platform.system())
        if not command:
            raise RuntimeError(f"Opening {key} is not configured for this operating system.")
        if platform.system() == "Linux" and shutil.which(command[0]) is None:
            raise RuntimeError(f"{command[0]} is not installed on this system.")
        subprocess.Popen(command, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, close_fds=True)
        return {"opened": key}

    def system_status(self) -> dict[str, Any]:
        import psutil

        battery = psutil.sensors_battery()
        try:
            temperatures = psutil.sensors_temperatures(fahrenheit=False)
            available_temperatures = [entry.current for group in temperatures.values() for entry in group if entry.current is not None]
        except (AttributeError, OSError):
            available_temperatures = []
        return {
            "cpu_percent": psutil.cpu_percent(interval=0.3),
            "memory_percent": psutil.virtual_memory().percent,
            "memory_used_gb": round(psutil.virtual_memory().used / (1024**3), 2),
            "memory_total_gb": round(psutil.virtual_memory().total / (1024**3), 2),
            "battery_percent": battery.percent if battery else None,
            "battery_plugged": battery.power_plugged if battery else None,
            "temperature_c": max(available_temperatures) if available_temperatures else None,
            "platform": platform.platform(),
        }

    def set_volume(self, percent: int) -> dict[str, int]:
        if not 0 <= percent <= 100:
            raise ValueError("Volume must be between 0 and 100 percent.")
        adapter = SystemSettingsAdapter()
        before = adapter.read_volume()
        adapter.set_volume(percent)
        after = adapter.read_volume()
        return {"before": before, "after": after}

    def web_search(self, query: str, limit: int = 6) -> list[dict[str, str]]:
        text = query.strip()
        if not text or len(text) > 500:
            raise ValueError("Search query must contain 1–500 characters.")
        count = max(1, min(limit, 10))
        try:
            from ddgs import DDGS

            results = DDGS().text(text, max_results=count)
        except Exception as exc:
            raise RuntimeError(f"Web search failed: {str(exc).splitlines()[0][:200]}") from exc
        return [
            {"title": str(item.get("title", "Untitled")), "url": str(item.get("href", "")), "snippet": str(item.get("body", ""))}
            for item in results if item.get("href")
        ]

    def browser_visit(self, url: str) -> dict[str, Any]:
        return BrowserController().visit(url)

    def browser_interact(self, url: str, action: str, selector: str, value: str = "") -> dict[str, Any]:
        return BrowserController().interact(url, action, selector, value)

    def get_weather(self, location: str) -> dict[str, Any]:
        place = location.strip()
        if not place or len(place) > 160:
            raise ValueError("Enter a city or postal code up to 160 characters.")
        geocode_url = "https://geocoding-api.open-meteo.com/v1/search?" + urlencode({"name": place, "count": 1, "language": "en", "format": "json"})
        geocoding = _http_json(geocode_url)
        results = geocoding.get("results") or []
        if not results:
            raise RuntimeError(f"No location match found for {place!r}.")
        found = results[0]
        params = {
            "latitude": found["latitude"],
            "longitude": found["longitude"],
            "current": "temperature_2m,relative_humidity_2m,apparent_temperature,is_day,precipitation,rain,weather_code,wind_speed_10m",
            "timezone": "auto",
        }
        forecast_url = "https://api.open-meteo.com/v1/forecast?" + urlencode(params)
        forecast = _http_json(forecast_url)
        current = forecast.get("current", {})
        weather_code = current.get("weather_code")
        return {
            "location": ", ".join(value for value in [found.get("name"), found.get("admin1"), found.get("country")] if value),
            "current": current,
            "condition": _WEATHER_DESCRIPTIONS.get(weather_code, "Unknown condition"),
            "units": forecast.get("current_units", {}),
            "source": forecast_url,
            "attribution": "Weather data by Open-Meteo; location data based on GeoNames.",
        }

    def file_search(self, pattern: str) -> list[str]:
        pattern = safe_pattern(pattern)
        matches: list[str] = []
        for root in self.approved_roots:
            base = Path(root).expanduser().resolve(strict=True)
            for path in base.rglob(pattern):
                if path.is_file():
                    matches.append(str(path))
                    if len(matches) >= 200:
                        return matches
        return matches

    def file_read(self, path: str) -> dict[str, str]:
        target = approved_path(path, self.approved_roots)
        if not target.is_file():
            raise FileNotFoundError("File does not exist.")
        data = target.read_bytes()
        suffix = target.suffix.lower()
        if suffix == ".pdf":
            if len(data) > MAX_DOCUMENT_BYTES:
                raise ValueError("PDFs are limited to 20 MB.")
            try:
                from pypdf import PdfReader
                reader = PdfReader(BytesIO(data))
                text = "\n\n".join((page.extract_text() or "") for page in reader.pages[:200])
            except Exception as exc:
                raise ValueError(f"PDF text extraction failed: {str(exc).splitlines()[0][:180]}") from exc
        elif suffix == ".docx":
            if len(data) > MAX_DOCUMENT_BYTES:
                raise ValueError("Word documents are limited to 20 MB.")
            try:
                from docx import Document
                document = Document(BytesIO(data))
                sections = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
                sections.extend(" | ".join(cell.text for cell in row.cells) for table in document.tables[:40] for row in table.rows[:100])
                text = "\n".join(sections)
            except Exception as exc:
                raise ValueError(f"Word document extraction failed: {str(exc).splitlines()[0][:180]}") from exc
        else:
            if len(data) > MAX_TEXT_FILE:
                raise ValueError("Text files are limited to 1 MB.")
            if suffix not in {".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".xml", ".py", ".js", ".ts", ".html", ".log", ".rst"}:
                raise ValueError("Supported local formats are UTF-8 text, PDF, and DOCX.")
            try:
                text = data.decode("utf-8")
            except UnicodeDecodeError as exc:
                raise ValueError("This file is not UTF-8 text.") from exc
        if len(text) > 400_000:
            text = text[:400_000] + "\n\n[Document text truncated at 400,000 characters.]"
        return {"path": str(target), "text": text}

    def file_write(self, path: str, content: str) -> dict[str, Any]:
        target = approved_path(path, self.approved_roots)
        if len(content.encode("utf-8")) > MAX_TEXT_FILE:
            raise ValueError("Writes are limited to 1 MB.")
        if not any(target.parent == Path(root).expanduser().resolve(strict=True) or Path(root).expanduser().resolve(strict=True) in target.parent.parents for root in self.approved_roots):
            raise ValueError("The parent folder must already exist in an approved folder.")
        existed = target.exists()
        if existed and not target.is_file():
            raise ValueError("Only regular files can be replaced.")
        original_stat = target.stat() if existed else None
        backup = target.read_bytes() if existed else b""
        if len(backup) > MAX_TEXT_FILE:
            raise ValueError("The existing file is larger than the 1 MB undo safety limit.")
        tmp = target.with_name(f".{target.name}.{uuid4().hex}.tmp")
        try:
            tmp.write_text(content, encoding="utf-8")
            if original_stat:
                os.chmod(tmp, stat.S_IMODE(original_stat.st_mode))
            tmp.replace(target)
        finally:
            tmp.unlink(missing_ok=True)
        written_stat = target.stat()
        payload = {
            "path": str(target), "existed": existed,
            "backup": base64.b64encode(backup).decode("ascii"),
            "written_sha256": _sha256(target.read_bytes()),
            "written_device": written_stat.st_dev,
            "written_inode": written_stat.st_ino,
            "original_mode": stat.S_IMODE(original_stat.st_mode) if original_stat else None,
            "original_atime_ns": original_stat.st_atime_ns if original_stat else None,
            "original_mtime_ns": original_stat.st_mtime_ns if original_stat else None,
        }
        return payload

    def file_move(self, source: str, destination: str) -> dict[str, Any]:
        src = approved_path(source, self.approved_roots)
        dst = approved_path(destination, self.approved_roots)
        if src.is_symlink() or not src.is_file():
            raise FileNotFoundError("Only existing regular files can be moved.")
        if dst.exists():
            raise FileExistsError("The destination already exists; no file was overwritten.")
        digest = _sha256(src.read_bytes())
        os.replace(src, dst)
        moved_stat = dst.stat()
        return {
            "source": str(src), "destination": str(dst), "sha256": digest,
            "device": moved_stat.st_dev, "inode": moved_stat.st_ino,
        }

    def file_delete(self, path: str) -> dict[str, Any]:
        source = approved_path(path, self.approved_roots)
        if source.is_symlink() or not source.is_file():
            raise FileNotFoundError("Only existing regular files can be moved to recoverable trash.")
        destination = source.with_name(f".{source.name}.{uuid4().hex}.zyvarnth-deleted")
        digest = _sha256(source.read_bytes())
        os.replace(source, destination)
        deleted_stat = destination.stat()
        return {
            "source": str(source), "trash": str(destination), "sha256": digest,
            "device": deleted_stat.st_dev, "inode": deleted_stat.st_ino,
        }

    def desktop_type_text(self, text: str) -> dict[str, Any]:
        if len(text) > 5000:
            raise ValueError("Typing is limited to 5,000 characters per request.")
        try:
            import pyautogui
        except ImportError as exc:
            raise RuntimeError("Desktop typing is optional. Install Zyvarnth with the desktop extra.") from exc
        pyautogui.write(text, interval=0.005)
        return {"typed_characters": len(text)}

    def desktop_shortcut(self, keys: list[Any]) -> dict[str, Any]:
        allowed = {"ctrl", "control", "command", "cmd", "alt", "option", "shift", "tab", "enter", "esc", "escape", "space", "up", "down", "left", "right", "home", "end", "delete", "backspace"}
        normalized = [str(key).lower() for key in keys]
        if not 1 <= len(normalized) <= 4 or any(key not in allowed for key in normalized):
            raise ValueError("Use one to four supported modifier/navigation keys.")
        try:
            import pyautogui
        except ImportError as exc:
            raise RuntimeError("Desktop shortcuts are optional. Install Zyvarnth with the desktop extra.") from exc
        pyautogui.hotkey(*normalized)
        return {"pressed": normalized}

    def create_reminder(self, title: str, minutes: int) -> dict[str, Any]:
        if not title.strip() or len(title) > 240 or not 1 <= minutes <= 525600:
            raise ValueError("Reminder title is required and time must be within 365 days.")
        due = datetime.now(timezone.utc) + timedelta(minutes=minutes)
        reminder_id = self.db.add_reminder(title.strip(), due.isoformat(timespec="seconds"))
        return {"id": reminder_id, "title": title.strip(), "due_at": due.isoformat(timespec="seconds")}

    def undo(self, action_id: int) -> str:
        entries = [action for action in self.db.actions(300) if action["id"] == action_id]
        if not entries:
            raise ValueError("Action not found.")
        item = entries[0]
        if item["undone"] or not item["reversible"]:
            raise ValueError("This action cannot be undone.")
        payload = item["payload"]
        name = item["tool_name"]
        if name == "file_write":
            target = Path(payload["path"])
            approved_path(target, self.approved_roots)
            current = target.stat() if target.exists() and not target.is_symlink() else None
            if (current is None or _sha256(target.read_bytes()) != payload["written_sha256"]
                    or current.st_dev != payload["written_device"] or current.st_ino != payload["written_inode"]):
                raise RuntimeError("The file changed after Zyvarnth wrote it; undo stopped to protect your edits.")
            if payload["existed"]:
                restored = target.with_name(f".{target.name}.{uuid4().hex}.undo")
                try:
                    restored.write_bytes(base64.b64decode(payload["backup"]))
                    os.chmod(restored, int(payload["original_mode"]))
                    os.utime(restored, ns=(int(payload["original_atime_ns"]), int(payload["original_mtime_ns"])))
                    restored.replace(target)
                finally:
                    restored.unlink(missing_ok=True)
            else:
                target.unlink()
        elif name == "file_move":
            source, destination = Path(payload["source"]), Path(payload["destination"])
            approved_path(source, self.approved_roots)
            approved_path(destination, self.approved_roots)
            current = destination.stat() if destination.exists() and not destination.is_symlink() else None
            if (source.exists() or current is None or _sha256(destination.read_bytes()) != payload["sha256"]
                    or current.st_dev != payload["device"] or current.st_ino != payload["inode"]):
                raise RuntimeError("A path changed after the move; undo stopped to prevent overwriting data.")
            os.replace(destination, source)
        elif name == "file_delete":
            source, trash = Path(payload["source"]), Path(payload["trash"])
            approved_path(source, self.approved_roots)
            approved_trash = approved_path(trash, self.approved_roots)
            if approved_trash.parent != source.parent:
                raise RuntimeError("The recoverable trash entry is outside its original folder; undo stopped.")
            current = trash.stat() if trash.exists() and not trash.is_symlink() else None
            if (source.exists() or current is None or _sha256(trash.read_bytes()) != payload["sha256"]
                    or current.st_dev != payload["device"] or current.st_ino != payload["inode"]):
                raise RuntimeError("The original path is occupied or the trashed file changed; undo stopped.")
            source.parent.mkdir(parents=True, exist_ok=True)
            os.replace(trash, source)
        elif name == "set_volume":
            adapter = SystemSettingsAdapter()
            current = adapter.read_volume()
            if current != int(payload["after"]):
                raise RuntimeError("Volume changed after Zyvarnth adjusted it; undo stopped to preserve your change.")
            adapter.set_volume(int(payload["before"]))
        else:
            raise ValueError("This action has no undo handler.")
        self.db.mark_undone(action_id)
        return f"Undid {item['summary']}"


def _sha256(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _http_json(url: str) -> dict[str, Any]:
    request = Request(url, headers={"User-Agent": "Zyvarnth/0.1 (desktop assistant)", "Accept": "application/json"})
    try:
        with urlopen(request, timeout=10) as response:
            result = json.loads(response.read(1_000_000))
    except Exception as exc:
        raise RuntimeError(f"Live weather request failed: {str(exc).splitlines()[0][:180]}") from exc
    if not isinstance(result, dict):
        raise RuntimeError("Weather service returned an unexpected response.")
    return result

